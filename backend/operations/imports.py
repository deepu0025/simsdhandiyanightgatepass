import csv
import io
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from zipfile import BadZipFile

from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.utils import timezone
from django.utils.dateparse import parse_date, parse_datetime
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from rest_framework.exceptions import ValidationError as APIValidationError

from .models import IMPORTED_VERIFIED_PAYMENT_STATUSES, ImportBatch, Registration, Ticket

IMPORT_COLUMNS = (
    "ticket_id",
    "ticket_number",
    "registration_id",
    "ticket_type",
    "name",
    "email",
    "phone",
    "amount_due",
    "payment_method",
    "upi_id",
    "transaction_id_utr",
    "payment_status",
    "verified_by",
    "verified_at",
    "ticket_status",
    "ticket_generated_at",
    "email_status",
    "email_sent_at",
    "scan_status",
    "checked_in_at",
    "checked_in_by",
    "created_at",
    "updated_at",
    "last_synced_at",
    "event_id",
    "event_name",
    "registration_created_at",
    "registration_desk",
)
REQUIRED_COLUMNS = {"ticket_id", "ticket_number", "ticket_type", "name", "email"}
PAYMENT_STATUSES = {"", "PENDING", "PAID", "CAPTURED", "VERIFIED", "REJECTED", "FAILED"}
TICKET_STATUSES = {"", "PENDING", "ISSUED", "CANCELLED", "CANCELED", "REVOKED"}
EMAIL_STATUSES = {"", "PENDING", "PROCESSING", "SENT", "FAILED"}
SCAN_STATUSES = {
    "",
    "PENDING",
    "CHECKED_IN",
    "SCANNED",
    "NOT_SCANNED",
    "NOT_CHECKED_IN",
    "ALREADY_CHECKED_IN",
}
DATETIME_COLUMNS = (
    "verified_at",
    "ticket_generated_at",
    "email_sent_at",
    "checked_in_at",
    "created_at",
    "updated_at",
    "last_synced_at",
    "registration_created_at",
)


def _read_csv(uploaded):
    try:
        text = uploaded.read().decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
    except (UnicodeDecodeError, csv.Error) as error:
        raise APIValidationError("The CSV file could not be read as UTF-8.") from error
    if not reader.fieldnames:
        raise APIValidationError("The uploaded CSV has no header row.")

    def rows():
        try:
            yield from enumerate(reader, start=2)
        except csv.Error as error:
            raise APIValidationError("The CSV contains malformed rows.") from error

    return list(reader.fieldnames), rows()


def _read_xlsx(uploaded):
    try:
        workbook = load_workbook(uploaded, read_only=True, data_only=True)
        worksheet = workbook.active
        values = worksheet.iter_rows(values_only=True)
        header_row = next(values, None)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as error:
        raise APIValidationError("The XLSX file could not be read.") from error
    if not header_row:
        raise APIValidationError("The uploaded XLSX has no header row.")
    headers = ["" if value is None else str(value).strip() for value in header_row]
    while headers and not headers[-1]:
        headers.pop()

    def rows():
        for line_number, values_row in enumerate(values, start=2):
            row = dict(zip(headers, values_row))
            extra_values = values_row[len(headers):]
            if any(value not in (None, "") for value in extra_values):
                row[None] = extra_values
            yield line_number, row

    return headers, rows(), workbook


def read_spreadsheet(uploaded):
    filename = uploaded.name.lower()
    workbook = None
    if filename.endswith(".csv"):
        headers, rows = _read_csv(uploaded)
    elif filename.endswith(".xlsx"):
        headers, rows, workbook = _read_xlsx(uploaded)
    else:
        raise APIValidationError("Choose a CSV or XLSX file.")

    normalized_headers = [header.strip().lower() for header in headers]
    if len(normalized_headers) != len(set(normalized_headers)):
        if workbook:
            workbook.close()
        raise APIValidationError("The spreadsheet contains duplicate column headings.")
    missing = sorted(REQUIRED_COLUMNS - set(normalized_headers))
    if missing:
        if workbook:
            workbook.close()
        raise APIValidationError(
            {"detail": "Required columns are missing.", "columns": missing}
        )
    unknown = sorted(set(normalized_headers) - set(IMPORT_COLUMNS))

    def cleaned_rows():
        try:
            for line_number, source_row in rows:
                row = {}
                for key, value in source_row.items():
                    column = (key or "").strip().lower()
                    if not column:
                        if value not in (None, "", []):
                            row["_unexpected_columns"] = "yes"
                        continue
                    if isinstance(value, (datetime, date)):
                        value = value.isoformat()
                    row[column] = "" if value is None else str(value).strip()
                extra_values = source_row.get(None, [])
                if any(value not in (None, "") for value in extra_values):
                    row["_unexpected_columns"] = "yes"
                yield line_number, row
        finally:
            if workbook:
                workbook.close()

    return unknown, cleaned_rows()


def _parse_timestamp(value):
    if not value:
        return None
    parsed = parse_datetime(value)
    if not parsed:
        parsed_date = parse_date(value)
        if parsed_date:
            parsed = datetime.combine(parsed_date, datetime.min.time())
    if not parsed:
        raise ValueError("must be an ISO date or datetime")
    if timezone.is_naive(parsed):
        parsed = timezone.make_aware(parsed)
    return parsed


def _registration_external_id(event_id, row):
    external = row.get("registration_id", "").strip()
    if external:
        return f"import:{event_id}:{external}"
    return f"import:{event_id}:ticket:{row['ticket_id']}"


def validate_row(line_number, row, event_id):
    errors = []
    if row.get("_unexpected_columns"):
        errors.append("row contains more cells than the header")
    for field in ("ticket_id", "ticket_number", "ticket_type", "name", "email"):
        if not row.get(field):
            errors.append(f"{field} is required")
    length_limits = {
        "ticket_id": 255,
        "ticket_number": 80,
        "registration_id": 210,
        "ticket_type": 120,
        "name": 160,
        "email": 254,
        "phone": 32,
        "payment_method": 40,
        "upi_id": 120,
        "transaction_id_utr": 120,
        "verified_by": 160,
    }
    for field, limit in length_limits.items():
        if len(row.get(field, "")) > limit:
            errors.append(f"{field} exceeds the {limit}-character limit")
    if row.get("email"):
        try:
            validate_email(row["email"])
        except ValidationError:
            errors.append("email is invalid")

    payment_status = row.get("payment_status", "").upper()
    ticket_status = row.get("ticket_status", "").upper()
    email_status = row.get("email_status", "").upper()
    scan_status = row.get("scan_status", "").upper()
    if payment_status not in PAYMENT_STATUSES:
        errors.append("payment_status is unsupported")
    if ticket_status not in TICKET_STATUSES:
        errors.append("ticket_status is unsupported")
    if email_status not in EMAIL_STATUSES:
        errors.append("email_status is unsupported")
    if scan_status not in SCAN_STATUSES:
        errors.append("scan_status is unsupported")

    amount = None
    if row.get("amount_due"):
        try:
            amount = Decimal(row["amount_due"])
            if not amount.is_finite() or amount < 0:
                raise InvalidOperation
            if amount.as_tuple().exponent < -2 or amount > Decimal("99999999.99"):
                raise InvalidOperation
            amount = amount.quantize(Decimal("0.01"))
        except (InvalidOperation, ValueError):
            errors.append("amount_due must be a non-negative decimal amount")

    timestamps = {}
    for field in DATETIME_COLUMNS:
        try:
            timestamps[field] = _parse_timestamp(row.get(field, ""))
        except ValueError:
            errors.append(f"{field} must be an ISO date or datetime")
            timestamps[field] = None

    if not row.get("ticket_id") or not row.get("ticket_number"):
        action = "REJECT"
    else:
        ticket = Ticket.objects.filter(external_ticket_id=row["ticket_id"]).first()
        if ticket is None:
            try:
                ticket = Ticket.objects.filter(pk=row["ticket_id"]).first()
            except (ValidationError, ValueError):
                ticket = None
        number_match = Ticket.objects.filter(ticket_number=row["ticket_number"]).first()
        if ticket and ticket.ticket_number != row["ticket_number"]:
            errors.append("ticket_id already belongs to a different ticket_number")
        if number_match and (ticket is None or number_match.pk != ticket.pk):
            errors.append("ticket_number already belongs to a different ticket")
        action = "UPDATE" if ticket else "CREATE"

    data = {
        **row,
        "_event_id": str(event_id),
        "_amount_due": str(amount) if amount is not None else "",
        "_timestamps": {key: value.isoformat() if value else "" for key, value in timestamps.items()},
        "_registration_external_id": _registration_external_id(event_id, row)
        if row.get("ticket_id")
        else "",
    }
    return {
        "row": line_number,
        "ticket_number": row.get("ticket_number", ""),
        "action": action if not errors else "REJECT",
        "errors": errors,
        "data": data,
    }


def create_import_batch(
    *,
    uploaded,
    event,
    actor,
    max_rows=10000,
    registration_source=Registration.Source.IMPORT,
):
    if uploaded is None or not getattr(uploaded, "name", ""):
        raise APIValidationError("Choose a CSV or XLSX file to import.")
    if uploaded.size > 10 * 1024 * 1024:
        raise APIValidationError("Upload a spreadsheet smaller than 10 MB.")
    unknown_columns, source_rows = read_spreadsheet(uploaded)
    staged_rows = []
    for line_number, row in source_rows:
        if line_number > max_rows + 1:
            raise APIValidationError(f"Spreadsheet files may contain at most {max_rows} data rows.")
        if not any(value.strip() for value in row.values()):
            staged_rows.append({
                "row": line_number,
                "ticket_number": "",
                "action": "SKIP",
                "errors": [],
                "data": {},
            })
            continue
        staged = validate_row(line_number, row, event.pk)
        if staged["data"]:
            staged["data"]["_registration_source"] = registration_source
        staged_rows.append(staged)

    seen_ids = set()
    seen_numbers = set()
    for row in staged_rows:
        data = row["data"]
        if not data:
            continue
        ticket_id = data.get("ticket_id", "")
        ticket_number = data.get("ticket_number", "")
        duplicates = []
        if ticket_id in seen_ids:
            duplicates.append("ticket_id is duplicated in this file")
        if ticket_number in seen_numbers:
            duplicates.append("ticket_number is duplicated in this file")
        if duplicates:
            row["errors"].extend(duplicates)
            row["action"] = "REJECT"
        if ticket_id:
            seen_ids.add(ticket_id)
        if ticket_number:
            seen_numbers.add(ticket_number)

    if not staged_rows:
        raise APIValidationError("The spreadsheet has no data rows.")

    batch = ImportBatch.objects.create(
        event=event,
        requested_by=actor,
        original_filename=uploaded.name[:255],
        rows=staged_rows,
        expires_at=timezone.now() + timedelta(hours=24),
    )
    return batch, unknown_columns


def batch_preview(batch, unknown_columns=None):
    created = sum(row["action"] == "CREATE" for row in batch.rows)
    updated = sum(row["action"] == "UPDATE" for row in batch.rows)
    rejected = sum(row["action"] == "REJECT" for row in batch.rows)
    skipped = sum(row["action"] == "SKIP" for row in batch.rows)
    email_evidence_rows = sum(
        row["action"] in {"CREATE", "UPDATE"}
        and not row["errors"]
        and str(row["data"].get("email_status", "")).strip().upper() == "SENT"
        and bool(row["data"].get("_timestamps", {}).get("email_sent_at"))
        for row in batch.rows
    )
    verified_payment_rows = sum(
        row["action"] in {"CREATE", "UPDATE"}
        and not row["errors"]
        and str(row["data"].get("payment_status", "")).strip().upper()
        in IMPORTED_VERIFIED_PAYMENT_STATUSES
        for row in batch.rows
    )
    return {
        "batch_id": str(batch.pk),
        "status": batch.status,
        "total_rows": len(batch.rows),
        "created": created,
        "updated": updated,
        "skipped": skipped,
        "rejected": rejected,
        "email_evidence_rows": email_evidence_rows,
        "verified_payment_rows": verified_payment_rows,
        "unknown_columns": unknown_columns or [],
        "rows": [
            {
                **{key: row[key] for key in ("row", "ticket_number", "action", "errors")},
                "email_delivery_evidence": (
                    row["action"] in {"CREATE", "UPDATE"}
                    and not row["errors"]
                    and str(row["data"].get("email_status", "")).strip().upper() == "SENT"
                    and bool(row["data"].get("_timestamps", {}).get("email_sent_at"))
                ),
                "payment_status": row["data"].get("payment_status", ""),
                "verified_payment": (
                    row["action"] in {"CREATE", "UPDATE"}
                    and not row["errors"]
                    and str(row["data"].get("payment_status", "")).strip().upper()
                    in IMPORTED_VERIFIED_PAYMENT_STATUSES
                ),
            }
            for row in batch.rows[:200]
        ],
        "errors": [
            {key: row[key] for key in ("row", "ticket_number", "action", "errors")}
            for row in batch.rows
            if row["errors"]
        ],
        "commit_available": (
            batch.status == ImportBatch.Status.PREVIEWED and (created + updated) > 0
        ),
        "expires_at": batch.expires_at,
        "result": batch.result,
    }
